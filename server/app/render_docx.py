"""Moteur de rendu Word déterministe (python-docx).

Le devis (JSON validé) + le profil d'entreprise → un .docx. Aucun remplacement de texte
dans un modèle existant : tout est construit ici, toujours de la même façon.
"""

from __future__ import annotations

import io
from typing import Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

from .config import Company
from .models import Quote
from .typo import NBSP, format_amount, format_date, format_percent, fr_text

TWIP_EMU = 635  # 1 twip = 635 EMU
PRICE_TABLE_WIDTHS = (6800, 2200)  # en twips : libellé / montant
DRAFT_MISSING = "BROUILLON — informations à compléter"
DRAFT_COMPANY = "BROUILLON — profil entreprise incomplet ou factice (SIRET, assurance…) : ne pas envoyer"

PRICE_PREFIX = {
    "complet": "Le prix forfaitaire total pour la fourniture, la main-d'œuvre et les prestations complètes, est de ",
    "main_oeuvre_seule": "Le prix forfaitaire total pour la main-d'œuvre est de ",
}
PRICE_SUFFIX = {
    "complet": "",
    "main_oeuvre_seule": " Les fournitures restent à la charge du client.",
}

# Ordre imposé par le schéma OOXML pour les enfants de <w:pPr> qui suivent <w:pBdr>.
_PPR_AFTER_BORDER = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


class RenderError(Exception):
    pass


def banner_reasons(quote: Quote, company: Company, draft: bool) -> list[str]:
    """Raisons d'afficher un bandeau BROUILLON dans l'en-tête."""
    reasons = []
    if not company.is_complete:
        # Garde-fou : jamais de devis « présentable » avec un SIRET vide ou factice.
        reasons.append(DRAFT_COMPANY)
    if draft and quote.missing:
        reasons.append(DRAFT_MISSING)
    return reasons


def price_amount_text(amount_ttc: int) -> str:
    """« 3 000 € TTC » d'un seul tenant (espaces insécables)."""
    return format_amount(amount_ttc) + NBSP + "TTC"


# --- Petits outils python-docx -------------------------------------------------------------


def _rgb(hex_color: str) -> RGBColor:
    try:
        return RGBColor.from_string(hex_color.strip().lstrip("#").upper())
    except ValueError:
        return RGBColor(0x2E, 0x6F, 0x8E)


def _missing(paragraph, label: str):
    """Champ manquant : texte entre crochets surligné en jaune."""
    run = paragraph.add_run(f"[{label}]")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def _text(paragraph, text: Optional[str], bold: bool = False, italic: bool = False):
    run = paragraph.add_run(fr_text(text or ""))
    run.bold = bold or None
    run.italic = italic or None
    return run


def _letter_spacing(run, twips: int) -> None:
    # À appeler juste après la création du run (avant toute autre mise en forme) pour respecter l'ordre OOXML.
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twips))
    run._r.get_or_add_rPr().append(spacing)


def _border(paragraph, side: str, color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), "6")
    edge.set(qn("w:space"), "4")
    edge.set(qn("w:color"), color)
    pbdr.append(edge)
    ppr.insert_element_before(pbdr, *_PPR_AFTER_BORDER)


def _field(paragraph, instruction: str) -> None:
    """Champ Word (PAGE, NUMPAGES…)."""
    def fld_char(kind: str):
        run = paragraph.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        run._r.append(el)

    fld_char("begin")
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    instr_run._r.append(instr)
    fld_char("separate")
    paragraph.add_run("1")
    fld_char("end")


def _set_cell_width(cell, emu: int) -> None:
    cell.width = Emu(emu)


# --- Construction du document -------------------------------------------------------------


class _DocxBuilder:
    def __init__(self, quote: Quote, company: Company, cover_photo: Optional[bytes], draft: bool):
        self.q = quote
        self.c = company
        self.cover_photo = cover_photo
        self.banner = banner_reasons(quote, company, draft)
        self.accent = _rgb(company.accent_color)
        self.doc = Document()

    # Mise en page et styles
    def setup(self) -> None:
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.54))
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)

        font = self.c.font or "American Typewriter"
        normal = self.doc.styles["Normal"]
        normal.font.name = font
        normal.font.size = Pt(10.5)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            rfonts.set(qn("w:eastAsia"), font)
            rfonts.set(qn("w:cs"), font)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "fr-FR")
        rpr.append(lang)
        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.line_spacing = 1.1

        heading = self.doc.styles.add_style("Devis Titre", WD_STYLE_TYPE.PARAGRAPH)
        heading.base_style = normal
        heading.font.bold = True
        heading.font.italic = True
        heading.font.underline = True
        heading.font.color.rgb = self.accent
        heading.font.size = Pt(12)
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(4)
        heading.paragraph_format.keep_with_next = True

        main = self.doc.styles.add_style("Devis Titre Principal", WD_STYLE_TYPE.PARAGRAPH)
        main.base_style = heading
        main.font.size = Pt(15)
        main.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        main.paragraph_format.space_before = Pt(14)
        main.paragraph_format.space_after = Pt(8)

        small = self.doc.styles.add_style("Devis Petit", WD_STYLE_TYPE.PARAGRAPH)
        small.base_style = normal
        small.font.size = Pt(7.5)
        small.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        small.paragraph_format.space_after = Pt(0)
        small.paragraph_format.line_spacing = 1.0

        props = self.doc.core_properties
        props.title = self.q.title or "Devis"
        props.author = self.c.label
        props.comments = ""

    def header_footer(self) -> None:
        section = self.doc.sections[0]
        header = section.header
        p = header.paragraphs[0]
        run = p.add_run(self.c.label)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = self.accent
        p.paragraph_format.space_after = Pt(0)
        if self.c.tagline:
            p2 = header.add_paragraph()
            r2 = p2.add_run(self.c.tagline)
            _letter_spacing(r2, 40)
            r2.font.size = Pt(8)
            r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            p2.paragraph_format.space_after = Pt(2)
            last = p2
        else:
            last = p
        _border(last, "bottom", str(self.accent))
        last.paragraph_format.space_after = Pt(6)
        for reason in self.banner:
            pb = header.add_paragraph()
            rb = pb.add_run(reason)
            rb.bold = True
            rb.font.size = Pt(8.5)
            rb.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            pb.paragraph_format.space_before = Pt(0)
            pb.paragraph_format.space_after = Pt(4)

        footer = section.footer
        c = self.c
        lines = [
            " — ".join(x for x in [c.label, c.legal_form, f"SIRET {c.siret}" if c.siret else "", f"NAF {c.naf}" if c.naf else ""] if x),
            c.address,
            " — ".join(x for x in [c.email, c.phone] if x),
            c.insurance_text,
        ]
        fp = footer.paragraphs[0]
        fp.style = self.doc.styles["Devis Petit"]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _border(fp, "top", "999999")
        first = True
        for line in lines:
            if not line:
                continue
            if not first:
                fp.add_run().add_break()
            fp.add_run(fr_text(line))
            first = False
        pn = footer.add_paragraph(style=self.doc.styles["Devis Petit"])
        pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pn.add_run("Page ")
        _field(pn, "PAGE")
        pn.add_run(" / ")
        _field(pn, "NUMPAGES")

    def cover(self) -> None:
        if not (self.q.options.cover_photo and self.cover_photo):
            return
        try:
            img = DocxImage.from_blob(self.cover_photo)
            ratio = img.px_height / img.px_width
        except Exception as exc:  # image illisible
            raise RenderError("La photo de la page de garde est illisible (JPEG ou PNG attendu).") from exc
        max_w, max_h = Cm(15.5), Cm(16)
        width = max_w
        height = int(width * ratio)
        if height > max_h:
            height = max_h
            width = int(height / ratio)
        p = self.doc.add_paragraph(style="Devis Titre Principal")
        p.paragraph_format.space_before = Pt(30)
        _text(p, self.q.title or "Devis")
        sub = self.doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _text(sub, f"Devis n° {self.q.number}" if self.q.number else "Devis")
        sub.add_run(" — " + format_date(self.q.date))
        pic = self.doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.space_before = Pt(18)
        pic.add_run().add_picture(io.BytesIO(self.cover_photo), width=Emu(width), height=Emu(height))
        pic.add_run().add_break(WD_BREAK.PAGE)

    def number_line(self) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run("Devis n°" + NBSP)
        run.bold = True
        if self.q.number:
            p.add_run(self.q.number).bold = True
        else:
            _missing(p, "Numéro attribué à la validation").bold = True
        p.add_run(" — " + format_date(self.q.date)).bold = True

    def parties(self) -> None:
        table = self.doc.add_table(rows=1, cols=2)
        table.autofit = False
        half = int(Cm(15.9) / 2)
        left, right = table.rows[0].cells
        _set_cell_width(left, half)
        _set_cell_width(right, half)

        def label(cell, text: str):
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = self.accent
            return p

        c = self.c
        label(left, "ENTREPRISE")
        _text(left.add_paragraph(), c.label, bold=True)
        for line in (c.address, c.phone, c.email):
            if line:
                _text(left.add_paragraph(), line)

        q = self.q
        label(right, "CLIENT")
        p = right.add_paragraph()
        if q.client.name:
            _text(p, q.client.name, bold=True)
        else:
            _missing(p, "Nom du client à compléter")
        p = right.add_paragraph()
        _text(p, "Chantier : ")
        if q.site_address:
            _text(p, q.site_address)
        else:
            _missing(p, "Adresse du chantier à compléter")
        for line in (q.client.phone, q.client.email):
            if line:
                _text(right.add_paragraph(), line)
        for cell in (left, right):
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(1)

    def title_and_intro(self) -> None:
        p = self.doc.add_paragraph(style="Devis Titre Principal")
        if self.q.title:
            _text(p, self.q.title)
        else:
            _missing(p, "Titre du devis à compléter")
        if self.q.intro:
            _text(self.doc.add_paragraph(), self.q.intro)

    def lots(self) -> None:
        if not self.q.lots:
            p = self.doc.add_paragraph()
            _missing(p, "Description des travaux à compléter")
            return
        for i, lot in enumerate(self.q.lots, 1):
            _text(self.doc.add_paragraph(style="Devis Titre"), f"{i}. {lot.title}")
            if lot.bullets:
                for bullet in lot.bullets:
                    bp = self.doc.add_paragraph(style="List Bullet")
                    bp.paragraph_format.space_after = Pt(2)
                    _text(bp, bullet)
            elif lot.paragraph:
                _text(self.doc.add_paragraph(), lot.paragraph)

    def price(self) -> None:
        q = self.q
        _text(self.doc.add_paragraph(style="Devis Titre"), "Prix")
        p = self.doc.add_paragraph()
        _text(p, PRICE_PREFIX[q.mode])
        if q.price.amount_ttc is not None:
            run = p.add_run(price_amount_text(q.price.amount_ttc))
            run.bold = True
        else:
            _missing(p, "Prix à définir").bold = True
        if q.price.vat_rate is not None:
            p.add_run(f" (TVA à {format_percent(q.price.vat_rate)}).")
        else:
            p.add_run(" (TVA" + NBSP + ": ")
            _missing(p, "Taux de TVA à définir")
            p.add_run(").")
        if PRICE_SUFFIX[q.mode]:
            _text(p, PRICE_SUFFIX[q.mode])
        if q.options.show_price_table:
            self.price_table()

    def price_table(self) -> None:
        q = self.q
        split = q.ht_and_vat()
        rows: list[tuple[str, Optional[str]]]
        if split is not None and q.price.amount_ttc is not None and q.price.vat_rate is not None:
            ht, vat = split
            rows = [
                ("Total HT", format_amount(ht, cents=True)),
                (f"TVA {format_percent(q.price.vat_rate)}", format_amount(vat, cents=True)),
                ("Total TTC", format_amount(q.price.amount_ttc, cents=True)),
            ]
        else:
            rows = [("Total HT", None), ("TVA", None), ("Total TTC", None)]
        table = self.doc.add_table(rows=0, cols=2)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for i, (label, amount) in enumerate(rows):
            cells = table.add_row().cells
            for cell, width in zip(cells, PRICE_TABLE_WIDTHS):
                _set_cell_width(cell, width * TWIP_EMU)
            last = i == len(rows) - 1
            cells[0].paragraphs[0].add_run(label).bold = last or None
            ap = cells[1].paragraphs[0]
            ap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if amount is None:
                _missing(ap, "Prix à définir")
            else:
                ap.add_run(amount).bold = last or None

    def conditions(self) -> None:
        q = self.q
        start = len(self.doc.paragraphs)
        _text(self.doc.add_paragraph(style="Devis Titre"), "Conditions")
        if q.payment_schedule:
            _text(self.doc.add_paragraph(), "Modalités de paiement :", bold=True)
            for step in q.payment_schedule:
                bp = self.doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)
                _text(bp, f"{step.label} : ")
                if step.amount_ttc is not None:
                    bp.add_run(price_amount_text(step.amount_ttc))
                else:
                    pct = step.percent or 0
                    bp.add_run(f"{format_percent(pct)} du montant TTC")
                    if q.price.amount_ttc is not None:
                        share = round(q.price.amount_ttc * pct / 100)
                        bp.add_run(", soit " + price_amount_text(share))
        p = self.doc.add_paragraph()
        _text(p, "Durée des travaux : ", bold=True)
        if q.duration:
            _text(p, q.duration)
        else:
            _missing(p, "Durée des travaux à compléter")
        if q.start_delay:
            p = self.doc.add_paragraph()
            _text(p, "Délai de démarrage : ", bold=True)
            _text(p, q.start_delay)
        _text(self.doc.add_paragraph(), f"Devis valable {q.validity_days} jours.")
        if q.nota:
            p = self.doc.add_paragraph()
            _text(p, "Nota : ", bold=True)
            _text(p, q.nota, italic=True)
        # Garde le bloc « Conditions » (et les signatures qui suivent) sur une même page.
        block = self.doc.paragraphs[start:]
        for para in block[:-1] if not q.options.show_signatures else block:
            para.paragraph_format.keep_with_next = True

    def signatures(self) -> None:
        if not self.q.options.show_signatures:
            return
        self.doc.add_paragraph().paragraph_format.keep_with_next = True
        table = self.doc.add_table(rows=1, cols=2)
        table.autofit = False
        for cell, text in zip(table.rows[0].cells, ("Signature entreprise :", "Signature client (bon pour accord) :")):
            _set_cell_width(cell, int(Cm(15.9) / 2))
            p = cell.paragraphs[0]
            p.paragraph_format.keep_with_next = True
            _text(p, text, bold=True)
            for _ in range(3):
                cell.add_paragraph()

    def build(self) -> bytes:
        self.setup()
        self.header_footer()
        self.cover()
        self.number_line()
        self.parties()
        self.title_and_intro()
        self.lots()
        self.price()
        self.conditions()
        self.signatures()
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()


def render_docx(quote: Quote, company: Company, cover_photo: Optional[bytes] = None, draft: bool = False) -> bytes:
    """Rend le devis en .docx (octets). `draft=True` ajoute le bandeau si des informations manquent."""
    if quote.company_id != company.id:
        raise RenderError("Le profil d'entreprise ne correspond pas au devis.")
    return _DocxBuilder(quote, company, cover_photo, draft).build()
