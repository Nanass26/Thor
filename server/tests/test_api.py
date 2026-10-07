import io
import logging
from datetime import date, timedelta

import pytest
from docx import Document

from app.config import Settings
from app.llm import LLMOutputError
from app.main import DailyQuota, QuotaExceeded, create_app, quote_filename
from app.models import Op, RespondOutput
from tests.conftest import AUTH, HAS_SOFFICE, b64, fake_jpeg, load_fixture, make_client, post_chat, tiny_png


# --- Santé et authentification ------------------------------------------------------------------


def test_health_sans_jeton_limite_a_ok(client):
    r = client.get("/v1/health")
    assert r.status_code == 200 and r.json() == {"ok": True}


def test_health_avec_jeton(client):
    body = client.get("/v1/health", headers=AUTH).json()
    assert body["ok"] is True and body["mode"] == "mock"
    assert isinstance(body["soffice"], bool) and body["version"]


def test_health_mauvais_jeton(client):
    r = client.get("/v1/health", headers={"Authorization": "Bearer faux"})
    assert r.status_code == 401
    assert "Jeton refusé" in r.json()["error"]


@pytest.mark.parametrize(
    "method,path",
    [("get", "/v1/companies"), ("post", "/v1/chat"), ("post", "/v1/extract"), ("post", "/v1/render/docx"), ("post", "/v1/render/pdf")],
)
def test_401_sans_jeton(client, method, path):
    r = getattr(client, method)(path, json={}) if method == "post" else client.get(path)
    assert r.status_code == 401
    assert "Jeton manquant" in r.json()["error"]


def test_entreprises_libelles_seulement(client):
    companies = client.get("/v1/companies", headers=AUTH).json()
    assert [c["id"] for c in companies] == ["levix", "sar_bat", "sb_concept", "m_bat", "id_bat"]
    assert all(set(c) == {"id", "label", "tagline"} for c in companies)  # rien de légal


def test_mode_live_refuse_de_demarrer_sans_jeton():
    with pytest.raises(RuntimeError, match="APP_TOKEN"):
        create_app(Settings(mock_llm=False, app_token="", anthropic_api_key="sk-test"))
    with pytest.raises(RuntimeError, match="ANTHROPIC_API_KEY"):
        create_app(Settings(mock_llm=False, app_token="x", anthropic_api_key=""))


def test_mode_demo_jeton_par_defaut(monkeypatch, tmp_path):
    monkeypatch.setenv("MOCK_LLM", "1")
    monkeypatch.delenv("APP_TOKEN", raising=False)
    settings = Settings.load(tmp_path / "absent.env")
    assert settings.mock_llm and settings.app_token == "demo"


def test_lecture_fichier_env(tmp_path, monkeypatch):
    for key in ("APP_TOKEN", "MOCK_LLM", "MAX_IMAGES"):
        monkeypatch.delenv(key, raising=False)
    env = tmp_path / ".env"
    env.write_text("APP_TOKEN=abc   # commentaire\nMAX_IMAGES=4\nMOCK_LLM=0\n# ligne ignorée\n", encoding="utf-8")
    s = Settings.load(env)
    assert s.app_token == "abc" and s.max_images == 4 and not s.mock_llm


# --- Limites ----------------------------------------------------------------------------------


def test_plafond_images(client):
    r = client.post(
        "/v1/chat",
        json={"message": {"text": "x", "images": [b64(fake_jpeg())] * 7}, "company_id": "levix"},
        headers=AUTH,
    )
    assert r.status_code == 400
    assert "Trop d'images" in r.json()["error"]


def test_images_heic_et_invalides_refusees(client):
    heic = b"\x00\x00\x00\x18ftypheic" + b"\x00" * 20
    for data, msg in ((heic, "HEIC"), (b"GIF89a....", "JPEG attendu")):
        r = client.post("/v1/chat", json={"message": {"text": "x", "images": [b64(data)]}, "company_id": "levix"}, headers=AUTH)
        assert r.status_code == 400 and msg in r.json()["error"]
    r = client.post("/v1/chat", json={"message": {"text": "x", "images": ["@@pas du base64@@"]}, "company_id": "levix"}, headers=AUTH)
    assert r.status_code == 400 and "base64" in r.json()["error"]


def test_png_et_data_url_acceptes(client):
    r = post_chat(client, "fais-moi un devis", images=[tiny_png()])
    assert r["client_sheet"] is not None
    body = {"message": {"text": "devis", "images": ["data:image/jpeg;base64," + b64(fake_jpeg())]}, "company_id": "levix"}
    assert client.post("/v1/chat", json=body, headers=AUTH).status_code == 200


def test_corps_trop_gros():
    small = make_client(max_body_mb=1)
    big = "A" * (2 * 1024 * 1024)
    r = small.post("/v1/chat", json={"message": {"text": "x", "images": [big]}, "company_id": "levix"}, headers=AUTH)
    assert r.status_code == 413
    assert "trop volumineuse" in r.json()["error"]


def test_plafond_quotidien():
    c = make_client(daily_call_cap=2)
    post_chat(c, "Peinture salon 2 800 TTC")  # 1 appel
    post_chat(c, "Peinture salon 2 800 TTC")  # 2 appels
    r = post_chat(c, "Peinture salon 2 800 TTC", expect=429)
    assert "Plafond quotidien" in r["error"]


def test_compteur_remis_a_zero_a_minuit():
    today = [date(2026, 10, 7)]
    quota = DailyQuota(1, today=lambda: today[0])
    quota.charge()
    with pytest.raises(QuotaExceeded):
        quota.charge()
    today[0] += timedelta(days=1)
    quota.charge()  # nouveau jour : compteur remis à zéro


def test_message_vide_et_entreprise_inconnue(client):
    assert "Message vide" in post_chat(client, "  ", expect=400)["error"]
    assert "Entreprise inconnue" in post_chat(client, "devis", company="inconnue", expect=400)["error"]


def test_erreur_de_validation_en_francais_sans_echo(client):
    r = client.post("/v1/chat", json={"message": {"text": "SECRET-123"}, "company_id": "levix", "intrus": 1}, headers=AUTH)
    assert r.status_code == 422
    body = r.json()
    assert body["error"].startswith("Requête invalide")
    assert "SECRET-123" not in r.text


# --- Journal sans contenu, erreurs masquées --------------------------------------------------------


def test_journal_sans_contenu(client, caplog):
    caplog.set_level(logging.INFO, logger="devischat")
    post_chat(client, "Peinture chez Mme SECRET-NOM, 2 800 TTC", images=[fake_jpeg()])
    assert "POST /v1/chat 200" in caplog.text
    assert "SECRET-NOM" not in caplog.text
    assert "Karim" not in caplog.text  # données lues dans la capture


class ExplodingLLM:
    def extract(self, text, images):
        raise RuntimeError("SECRET-TRACE contenu client")

    def respond(self, ctx):
        raise RuntimeError("SECRET-TRACE contenu client")


def test_erreur_interne_masquee(caplog):
    caplog.set_level(logging.INFO, logger="devischat")
    c = make_client(llm=ExplodingLLM())
    r = post_chat(c, "devis", expect=500)
    assert r == {"error": "Erreur interne du serveur."}
    assert "SECRET-TRACE" not in caplog.text


# --- Relance unique en cas de réponse invalide -------------------------------------------------------


class ScriptedLLM:
    """Renvoie successivement les sorties prévues (ou lève les erreurs prévues)."""

    def __init__(self, *outputs):
        self.outputs = list(outputs)
        self.calls = 0
        self.contexts = []

    def extract(self, text, images):
        raise AssertionError("pas d'image dans ces tests")

    def respond(self, ctx):
        self.calls += 1
        self.contexts.append(ctx.previous_error)
        out = self.outputs.pop(0)
        if isinstance(out, Exception):
            raise out
        return out


VALID_PRICE = RespondOutput(intent="modify", assistant_text="ok", ops=[Op(op="replace", path="/price/amount_ttc", value=3500)])


def test_une_seule_relance_puis_succes():
    llm = ScriptedLLM(LLMOutputError(["champ manquant"]), VALID_PRICE)
    r = post_chat(make_client(llm=llm), "mets 3500 TTC", quote=load_fixture("mo_seule_appartement_53m2"))
    assert r["quote"]["price"]["amount_ttc"] == 3500
    assert llm.calls == 2 and llm.contexts == [None, "champ manquant"]


def test_deux_reponses_invalides_erreur_propre():
    llm = ScriptedLLM(LLMOutputError(["x"]), LLMOutputError(["y"]))
    r = post_chat(make_client(llm=llm), "mets 3500 TTC", quote=load_fixture("mo_seule_appartement_53m2"), expect=502)
    assert "deux réponses invalides" in r["error"]
    assert llm.calls == 2


def test_operations_refusees_deux_fois_rien_n_est_change():
    bad = RespondOutput(intent="modify", assistant_text="ok", ops=[Op(op="replace", path="/number", value="99/26")])
    llm = ScriptedLLM(bad, bad)
    r = post_chat(make_client(llm=llm), "change le numéro", quote=load_fixture("mo_seule_appartement_53m2"))
    assert r["intent"] == "question" and r["quote"] is None
    assert "Rien n'a été changé" in r["assistant_text"]
    assert "interdit" in llm.contexts[1]  # l'erreur a bien été renvoyée au modèle


def test_changement_de_mode_remet_le_prix_a_definir():
    only_mode = RespondOutput(intent="modify", assistant_text="ok", ops=[Op(op="replace", path="/mode", value="complet")])
    r = post_chat(make_client(llm=ScriptedLLM(only_mode)), "version complète", quote=load_fixture("mo_seule_appartement_53m2"))
    assert r["quote"]["mode"] == "complet" and r["quote"]["price"]["amount_ttc"] is None


# --- Rendu via l'API ---------------------------------------------------------------------------


def test_render_docx(client):
    q = load_fixture("fenetre_bois")
    r = client.post("/v1/render/docx?draft=1", json={"quote": q.model_dump(mode="json")}, headers=AUTH)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="Devis_SAR_BAT_M_B_Exemple_Remplacement_d_une_fenetre_bois.docx"' in r.headers["content-disposition"]
    doc = Document(io.BytesIO(r.content))
    assert any("Devis valable 30 jours." in p.text for p in doc.paragraphs)


def test_render_page_de_garde_et_entreprise_inconnue(client):
    q = load_fixture("fenetre_bois").model_dump(mode="json")
    q["options"]["cover_photo"] = True
    r = client.post("/v1/render/docx", json={"quote": q, "cover_photo_b64": b64(tiny_png())}, headers=AUTH)
    assert r.status_code == 200 and len(Document(io.BytesIO(r.content)).inline_shapes) == 1
    q["company_id"] = "inconnue"
    r = client.post("/v1/render/docx", json={"quote": q}, headers=AUTH)
    assert r.status_code == 400 and "Entreprise inconnue" in r.json()["error"]


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice absent")
def test_render_pdf(client):
    q = load_fixture("mo_seule_appartement_53m2")
    r = client.post("/v1/render/pdf?draft=1", json={"quote": q.model_dump(mode="json")}, headers=AUTH)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert r.headers["content-type"] == "application/pdf"


def test_render_pdf_sans_libreoffice(monkeypatch):
    import app.main as main_module

    monkeypatch.setattr(main_module, "find_soffice", lambda configured="": None)
    c = make_client()
    q = load_fixture("mo_seule_appartement_53m2")
    r = c.post("/v1/render/pdf", json={"quote": q.model_dump(mode="json")}, headers=AUTH)
    assert r.status_code == 503 and r.json()["error"] == "LibreOffice n'est pas installé sur le serveur."
    assert c.get("/v1/health", headers=AUTH).json()["soffice"] is False


def test_nom_de_fichier_nettoye(companies):
    q = load_fixture("optique_complet_55m2")
    q.client.name = 'A/B:C*D?"E<F>|G'
    name = quote_filename(q, companies["sb_concept"], "pdf")
    assert name.startswith("Devis_SB_CONCEPT_A_B_C_D_E_F_G_")
    assert not any(ch in name for ch in '/\\:*?"<>|')
