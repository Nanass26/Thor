import io
import subprocess

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from app.config import Company
from app.models import Quote
from app.render_docx import DRAFT_COMPANY, DRAFT_MISSING, RenderError, banner_reasons, price_amount_text, render_docx
from app.render_pdf import PdfError, docx_to_pdf
from app.typo import NBSP, plain_spaces
from tests.conftest import FIXTURE_NAMES, HAS_POPPLER, HAS_SOFFICE, load_fixture, tiny_png

REAL_COMPANY = Company(
    id="levix", label="LEVIX", tagline="DESIGN & BUILD", legal_form="SAS", siret="123 456 789 00012",
    naf="4334Z", address="10 rue Imaginaire, 92400 Courbevoie", email="contact@levix.fr",
    phone="01 23 45 67 89", insurance_text="Assurance décennale n° 123 — Assureur Imaginaire",
)


def open_docx(data: bytes):
    return Document(io.BytesIO(data))


def body_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts += [p.text for p in cell.paragraphs]
    return "\n".join(parts)


def header_text(doc) -> str:
    return "\n".join(p.text for p in doc.sections[0].header.paragraphs)


def all_runs(doc):
    for p in doc.paragraphs:
        yield from p.runs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield from p.runs


def test_docx_contenu_de_reference(ref_quote, companies):
    doc = open_docx(render_docx(ref_quote, companies["levix"]))
    text = plain_spaces(body_text(doc))
    assert "Le prix forfaitaire total pour la main-d’œuvre est de 3 000 € TTC (TVA à 10 %). Les fournitures restent à la charge du client." in text
    assert "Devis n° [Numéro attribué à la validation] — 7 octobre 2026" in text
    assert "Devis valable 30 jours." in text
    assert "Durée des travaux : 3 semaines" in text
    assert "1. Protection et préparation" in text and "3. Sols" in text


def test_montant_d_un_seul_tenant(ref_quote, companies):
    doc = open_docx(render_docx(ref_quote, companies["levix"]))
    amount = price_amount_text(3000)
    assert amount == f"3 000{NBSP}€{NBSP}TTC"
    runs = [r for r in all_runs(doc) if r.text == amount]
    assert len(runs) == 1 and runs[0].bold


def test_phrase_prix_complet_et_numero(companies):
    q = load_fixture("optique_complet_55m2")
    doc = open_docx(render_docx(q, companies["sb_concept"]))
    text = plain_spaces(body_text(doc))
    assert "Le prix forfaitaire total pour la fourniture, la main-d’œuvre et les prestations complètes, est de 6 500 € TTC (TVA à 20 %)." in text
    assert "Devis n° 18/26 — 14 septembre 2026" in text
    # Tableau de prix demandé : HT / TVA / TTC, colonnes 6 800 / 2 200 twips
    assert "Total HT" in text and "5 416,67 €" in text and "1 083,33 €" in text and "6 500,00 €" in text
    price_table = doc.tables[1]
    assert [c.width for c in price_table.rows[0].cells] == [6800 * 635, 2200 * 635]
    assert "Signature client (bon pour accord) :" in text


def test_tableau_et_signatures_absents_par_defaut(ref_quote, companies):
    doc = open_docx(render_docx(ref_quote, companies["levix"]))
    text = body_text(doc)
    assert "Total HT" not in text
    assert "Signature entreprise" not in text
    assert len(doc.tables) == 1  # seulement le bloc entreprise / client


def test_crochets_jaunes_sans_prix(companies):
    q = load_fixture("incomplet_sans_prix")
    doc = open_docx(render_docx(q, companies["id_bat"]))
    text = body_text(doc)
    assert "[Prix à définir]" in text
    assert "[Nom du client à compléter]" in text
    assert "[Adresse du chantier à compléter]" in text
    assert "€ TTC" not in plain_spaces(text)  # aucun montant inventé
    yellow = [r.text for r in all_runs(doc) if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert "[Prix à définir]" in yellow and "[Nom du client à compléter]" in yellow


def test_echeancier(companies):
    q = load_fixture("peinture_acompte")
    text = plain_spaces(body_text(open_docx(render_docx(q, companies["m_bat"]))))
    assert "Acompte à la signature : 2 500 € TTC" in text
    assert "Solde à la fin des travaux : 2 000 € TTC" in text


def test_echeancier_en_pourcentage(ref_quote, companies):
    data = ref_quote.model_dump(mode="json")
    data["payment_schedule"] = [{"label": "Acompte", "percent": 30}]
    q = Quote.model_validate(data)
    text = plain_spaces(body_text(open_docx(render_docx(q, companies["levix"]))))
    assert "Acompte : 30 % du montant TTC, soit 900 € TTC" in text


def test_pied_de_page_vient_du_profil(ref_quote):
    doc = open_docx(render_docx(ref_quote, REAL_COMPANY))
    footer = plain_spaces("\n".join(p.text for p in doc.sections[0].footer.paragraphs))
    for value in ("SIRET 123 456 789 00012", "NAF 4334Z", "10 rue Imaginaire", "contact@levix.fr", "Assurance décennale n° 123"):
        assert value in footer
    assert header_text(doc).startswith("LEVIX\nDESIGN & BUILD")


def test_style_maison(ref_quote, companies):
    doc = open_docx(render_docx(ref_quote, companies["levix"]))
    section = doc.sections[0]
    assert round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7
    assert round(section.left_margin.cm, 2) == 2.54 and round(section.top_margin.cm, 2) == 2.54
    assert doc.styles["Normal"].font.name == "American Typewriter"
    heading = doc.styles["Devis Titre"].font
    assert heading.bold and heading.italic and heading.underline
    assert str(heading.color.rgb) == "2E6F8E"


def test_bandeau_brouillon(ref_quote, companies):
    # Profil factice → bandeau toujours présent (garde-fou SIRET)
    assert banner_reasons(ref_quote, companies["levix"], draft=False) == [DRAFT_COMPANY]
    # Vrai profil : bandeau « à compléter » seulement avec draft=1 et des champs manquants
    assert banner_reasons(ref_quote, REAL_COMPANY, draft=False) == []
    assert banner_reasons(ref_quote, REAL_COMPANY, draft=True) == [DRAFT_MISSING]
    complete = load_fixture("fenetre_bois")
    real = Company(**{**REAL_COMPANY.__dict__, "id": "sar_bat"})
    assert banner_reasons(complete, real, draft=True) == []
    assert DRAFT_MISSING in header_text(open_docx(render_docx(ref_quote, REAL_COMPANY, draft=True)))


def test_profil_factice_detecte(companies):
    for company in companies.values():
        assert not company.is_complete
    assert REAL_COMPANY.is_complete


def test_page_de_garde(ref_quote, companies):
    q = ref_quote.model_copy(update={"options": ref_quote.options.model_copy(update={"cover_photo": True})})
    with_cover = open_docx(render_docx(q, companies["levix"], cover_photo=tiny_png(30, 60)))
    assert len(with_cover.inline_shapes) == 1
    assert with_cover.inline_shapes[0].height.cm <= 16.01
    without_option = open_docx(render_docx(ref_quote, companies["levix"], cover_photo=tiny_png()))
    assert len(without_option.inline_shapes) == 0
    with pytest.raises(RenderError):
        render_docx(q, companies["levix"], cover_photo=b"pas une image")


def test_entreprise_incoherente(ref_quote, companies):
    with pytest.raises(RenderError):
        render_docx(ref_quote, companies["sar_bat"])


# --- PDF ------------------------------------------------------------------------------------


def pdf_pages_and_text(pdf: bytes, tmp_path) -> tuple[int, str]:
    path = tmp_path / "devis.pdf"
    path.write_bytes(pdf)
    info = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True, check=True).stdout
    pages = int(next(line.split()[-1] for line in info.splitlines() if line.startswith("Pages:")))
    text = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, check=True).stdout
    return pages, plain_spaces(text)


@pytest.mark.skipif(not (HAS_SOFFICE and HAS_POPPLER), reason="LibreOffice ou poppler absent")
@pytest.mark.parametrize("name", FIXTURE_NAMES)
def test_pdf_une_a_trois_pages_et_montant(name, companies, tmp_path):
    q = load_fixture(name)
    pdf = docx_to_pdf(render_docx(q, companies[q.company_id], draft=True))
    pages, text = pdf_pages_and_text(pdf, tmp_path)
    assert 1 <= pages <= 3
    if q.price.amount_ttc is not None:
        expected = plain_spaces(price_amount_text(q.price.amount_ttc))
        assert expected in text
    else:
        assert "[Prix à définir]" in text


def test_pdf_sans_libreoffice(monkeypatch):
    import app.render_pdf as rp

    monkeypatch.setattr(rp, "find_soffice", lambda configured="": None)
    with pytest.raises(PdfError) as exc:
        rp.docx_to_pdf(b"x")
    assert "LibreOffice n'est pas installé" in str(exc.value)
